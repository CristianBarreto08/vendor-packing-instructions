"""Re-open the saved workbook and check it -- not running without errors is not enough.

What gets sent is an Excel file, and the ways it can go wrong *without* an
error are exactly the ones that matter:

- a formula still pointing at another workbook (what this tool removes)
- an Excel error value (#REF!, #NAME?, #N/A) in a cell
- the table header overwritten by the first product
- leftover template rows under the data, or under the invoice total
- a TOTAL PIECES sum that differs from the plan
- a row without the PO number
- missing, altered, or unanchored sample images; a missing template tab
"""

from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from . import media, template
from .assemble import BuildResult
from .columns import BOX, BOX_SHEET, INVOICE_SHEET, INVOICE_SIGNATURE, LABELS, LABELS_SHEET
from .inputs import sheet_names

_EXTERNAL = re.compile(r"\[\d+\]")
_ERROR = re.compile(r"#(REF!|NAME\?|VALUE!|DIV/0!|N/A)")


def check_output(output: Path, result: BuildResult, po: int | None,
                 template_path: Path) -> list[str]:
    """Problems found in the saved file. Empty list = all good."""
    problems = _check_images(template_path, output)
    wb = openpyxl.load_workbook(output)
    try:
        box, labels = wb[BOX_SHEET], wb[LABELS_SHEET]
        for ws in (box, labels):
            _no_external_refs_or_errors(ws, problems)

        box_table = _header_still_there(box, template.box_table, problems)
        labels_table = _header_still_there(labels, template.labels_table, problems)
        if box_table:
            _no_leftovers(box, box_table.first_data + result.row_count, BOX["item_sku"], problems)
            _totals_match(box, box_table.first_data, result, problems)
        if labels_table:
            _no_leftovers(labels, labels_table.first_data + result.row_count,
                          LABELS["sku"], problems)
        _po_on_every_row(box, result, po, problems)
        _invoice_list_clean(wb[INVOICE_SHEET], result, problems)

        expected_tabs = sheet_names(template_path)
        missing_tabs = [t for t in expected_tabs if t not in wb.sheetnames]
        if missing_tabs:
            problems.append("missing template tabs: " + ", ".join(missing_tabs))
    finally:
        wb.close()
    return problems


def _check_images(template_path: Path, output: Path) -> list[str]:
    before, after = media.image_hashes(template_path), media.image_hashes(output)
    missing = sorted(set(before) - set(after))
    if missing:
        return [f"{len(missing)} template image(s) missing ({', '.join(missing)}): "
                "the vendor would receive blank label-sample tabs"]
    changed = sorted(n for n in before if before[n] != after.get(n))
    if changed:
        return [f"template image(s) altered or truncated: {', '.join(changed)}"]
    if media.anchor_count(template_path) != media.anchor_count(output):
        return ["images are present but their anchors do not match the template"]
    return []


def _no_external_refs_or_errors(ws: Worksheet, problems: list[str]) -> None:
    external, errors = [], []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                if cell.value.startswith("=") and _EXTERNAL.search(cell.value):
                    external.append(cell.coordinate)
                if _ERROR.search(cell.value):
                    errors.append(cell.coordinate)
    if external:
        problems.append(f"{ws.title}: {len(external)} formula(s) still point to another "
                        f"workbook ({', '.join(external[:5])})")
    if errors:
        problems.append(f"{ws.title}: {len(errors)} cell(s) with Excel errors "
                        f"({', '.join(errors[:5])})")


def _header_still_there(ws: Worksheet, locate, problems: list[str]):
    try:
        return locate(ws)
    except template.TemplateChanged as exc:
        problems.append(f"{ws.title}: table header is gone from the saved file. {exc}")
        return None


def _no_leftovers(ws: Worksheet, first_free: int, column: int, problems: list[str]) -> None:
    leftovers = [ws.cell(r, column).coordinate for r in range(first_free, ws.max_row + 1)
                 if ws.cell(r, column).value not in (None, "")]
    if leftovers:
        problems.append(f"{ws.title}: old template rows below the data "
                        f"({', '.join(leftovers[:5])})")


def _totals_match(ws: Worksheet, first: int, result: BuildResult, problems: list[str]) -> None:
    """Sum column D (values; openpyxl cannot evaluate E) against the plan."""
    total = sum(v for r in range(first, ws.max_row + 1)
                if isinstance(v := ws.cell(r, BOX["total_pcs"]).value, (int, float)))
    if abs(total - result.total_pieces) > 0.5:
        problems.append(f"{ws.title}: TOTAL PIECES sums to {total:,.0f}, "
                        f"expected {result.total_pieces:,.0f}")


def _po_on_every_row(ws: Worksheet, result: BuildResult, po: int | None,
                     problems: list[str]) -> None:
    if po is None:
        return
    missing = sum(1 for g in result.groups for p in g.packs
                  if ws.cell(p.box_row, BOX["po"]).value != po)
    if missing:
        problems.append(f"{ws.title}: {missing} row(s) without PO #{po}")


def _invoice_list_clean(ws: Worksheet, result: BuildResult, problems: list[str]) -> None:
    header = template.find_row(ws, INVOICE_SIGNATURE)
    if header is None:
        problems.append(f"{INVOICE_SHEET}: product list header not found")
        return
    last = header + len(result.groups) + 1  # + TOTAL row
    leftovers = [ws.cell(r, 1).coordinate for r in range(last + 1, ws.max_row + 1)
                 if ws.cell(r, 1).value not in (None, "")]
    if leftovers:
        problems.append(f"{INVOICE_SHEET}: example rows left under the total "
                        f"({', '.join(leftovers[:5])})")
