"""Find the input files and read the vendor tab of the packing plan.

Files are recognized by *content*, never by name: exports arrive with names
like ``500123.xlsx`` or ``Packing Plan 2026-09-10 - Vendor A PO 2041.xlsx``.
The only thing taken from a filename is the PO number.

Workspace layout
----------------
``Input/``      what the operator drops in: the approved packing plan and
                the sales report (priority column).
``Inbox/``      what the automated ERP export downloads: case dimensions.
``Reference/``  slow-changing files: vendor template, marketplace SKU map,
                barcode list.

``Input/`` and ``Inbox/`` are separate on purpose. Dimensions are no longer
something a person fetches; if an old dimensions export is left lying around
in ``Input/`` it is ignored (and flagged) instead of silently building the
vendor file with box sizes from a previous PO.
"""

from __future__ import annotations

import logging
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from .columns import (
    DIMENSIONS_COLUMNS,
    PLAN,
    PLAN_FIXED_TABS,
    PLAN_LAST_COL,
    PLAN_SIGNATURE,
    SALES_SHEET,
)

log = logging.getLogger(__name__)

_EXCEL_ERRORS = ("#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?")


class InputError(Exception):
    """A required file is missing, or the file present is not what we expect."""


def xlsx_files(folder: Path) -> list[Path]:
    """Real workbooks in a folder, skipping Excel's ``~$`` lock files."""
    if not folder.is_dir():
        return []
    return sorted(p for p in folder.glob("*.xlsx") if not p.name.startswith("~$"))


def sheet_names(path: Path) -> list[str]:
    """Sheet names read straight from the zip.

    Opening a 20 MB reference workbook with openpyxl just to list its tabs
    takes minutes; reading ``xl/workbook.xml`` takes milliseconds.
    """
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("xl/workbook.xml").decode("utf8", "replace")
    except (zipfile.BadZipFile, KeyError, OSError):
        return []
    return [n.replace("&amp;", "&") for n in re.findall(r'<sheet[^>]*name="([^"]*)"', xml)]


def _first_row(path: Path, sheet: str, max_col: int = 30) -> list:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        row = next(wb[sheet].iter_rows(min_row=1, max_row=1, max_col=max_col,
                                       values_only=True), ())
        return list(row)
    finally:
        wb.close()


def classify(folder: Path) -> dict[str, Path]:
    """Map ``{"plan" | "sales" | "dimensions": path}`` for one folder.

    Files are visited oldest to newest, so if two dimension exports from
    different runs are present, the most recent download wins rather than
    whichever sorts first alphabetically.
    """
    found: dict[str, Path] = {}
    for path in sorted(xlsx_files(folder), key=lambda p: p.stat().st_mtime):
        sheets = sheet_names(path)
        lower = {s.lower() for s in sheets}
        if "uploadtemplate" in lower and vendor_tabs(path, sheets):
            found["plan"] = path
        elif SALES_SHEET.lower() in lower:
            found["sales"] = path
        elif len(sheets) == 1:
            header = [str(v or "").strip() for v in _first_row(path, sheets[0])]
            if all(c in header for c in DIMENSIONS_COLUMNS[:2]):
                found["dimensions"] = path
    return found


def vendor_tabs(path: Path, sheets: list[str] | None = None) -> list[str]:
    """Tabs of the packing plan whose header matches the vendor-tab signature.

    Decided by header, not by name: vendors change, and a fixed list of tab
    names would silently skip a new one.
    """
    sheets = sheets or sheet_names(path)
    candidates = [s for s in sheets if s.lower() not in PLAN_FIXED_TABS]
    if not candidates:
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        tabs = []
        for name in candidates:
            row = next(wb[name].iter_rows(min_row=1, max_row=1, max_col=20,
                                          values_only=True), None)
            if row and _matches_signature(row):
                tabs.append(name)
        return tabs
    finally:
        wb.close()


def _matches_signature(row: tuple) -> bool:
    return all(i < len(row) and str(row[i] or "").strip().lower() == title.lower()
               for i, title in PLAN_SIGNATURE.items())


def po_from_filename(path: Path) -> int | None:
    """PO number embedded in a filename, or None.

    ``"Packing Plan 2026-09-10 - Vendor A po 2041.xlsx"`` -> ``2041``.
    No number is ever invented: when the name has none, the PO # column is
    left blank for the operator to fill in.
    """
    m = re.search(r"\bp\.?o\.?\s*#?\s*(\d{2,6})\b", path.stem, re.IGNORECASE)
    return int(m.group(1)) if m else None


@dataclass
class PlanRow:
    """One row of the vendor tab, already interpreted."""

    row: int
    sku: str
    product: str
    parent: str
    case_packs: float | None
    pack: float | None
    ordered: float | None
    cases: float | None
    total: float | None

    @property
    def is_child(self) -> bool:
        return bool(self.parent)


def _text(values: list, i: int) -> str:
    if i >= len(values) or values[i] is None:
        return ""
    text = str(values[i]).strip()
    return "" if text in _EXCEL_ERRORS else text


def _number(values: list, i: int) -> float | None:
    """Numeric value, or None when blank or an Excel error.

    ``#N/A`` is deliberately *not* turned into zero: a zero looks like valid
    data and would ship a zero-quantity line to the vendor; a None is visible
    and gets flagged in the review.
    """
    if i >= len(values):
        return None
    value = values[i]
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, str):
        value = value.strip()
        if not value or value.startswith("#"):
            return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_plan_tab(path: Path, tab: str) -> list[PlanRow]:
    """All rows with a SKU in the vendor tab, in sheet order."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if tab not in wb.sheetnames:
            raise InputError(f'workbook has no tab "{tab}" (tabs: {", ".join(wb.sheetnames)})')
        rows = []
        for n, raw in enumerate(wb[tab].iter_rows(min_row=2, max_col=PLAN_LAST_COL,
                                                  values_only=True), start=2):
            values = list(raw)
            sku = _text(values, PLAN["sku"])
            if not sku:
                continue
            rows.append(PlanRow(
                row=n, sku=sku,
                product=_text(values, PLAN["product"]),
                parent=_text(values, PLAN["parent"]),
                case_packs=_number(values, PLAN["case_packs"]),
                pack=_number(values, PLAN["pack"]),
                ordered=_number(values, PLAN["ordered"]),
                cases=_number(values, PLAN["cases"]),
                total=_number(values, PLAN["total"]),
            ))
        log.info("Read %d rows from tab %r", len(rows), tab)
        return rows
    finally:
        wb.close()
