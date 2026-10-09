"""The four lookup tables that translate the packing plan into vendor language.

    Marketplace SKU map   internal child SKU -> marketplace SKU
    Barcode list          marketplace SKU    -> barcode + label name
    ERP catalog export    internal child SKU -> case dimensions + packing notes
    Sales report          parent SKU         -> priority

Every lookup returns ``None`` for a miss, never a default. Whatever is not
found must show up highlighted in the workbook and listed in REVIEW.txt, not
be quietly filled in.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .columns import (
    BARCODE,
    BARCODE_SHEET,
    MARKETPLACE,
    MARKETPLACE_SHEET,
    SALES_PRIORITY_HEADER,
    SALES_SHEET,
    SALES_SKU_HEADER,
    normalize,
)

log = logging.getLogger(__name__)


def _clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.startswith("#") else text


def _at(row: tuple, i: int | None) -> object:
    return row[i] if i is not None and i < len(row) else None


# ---------------------------------------------------------------------------
# Marketplace SKU map
# ---------------------------------------------------------------------------

def load_marketplace_skus(path: Path) -> dict[str, str | None]:
    """``{internal SKU: marketplace SKU}``. First occurrence wins.

    "First wins" mirrors the spreadsheet lookup this replaces, so results are
    identical to the manual process even when the map has duplicate keys.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        mapping: dict[str, str | None] = {}
        for row in wb[MARKETPLACE_SHEET].iter_rows(min_row=2, max_col=5, values_only=True):
            internal = _clean(_at(row, MARKETPLACE["internal_sku"]))
            if internal and internal not in mapping:
                mapping[internal] = _clean(_at(row, MARKETPLACE["marketplace_sku"])) or None
        return mapping
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Barcode list
# ---------------------------------------------------------------------------

@dataclass
class BarcodeCatalog:
    """The barcode list indexed by both keys it can be searched with."""

    by_marketplace: dict[str, tuple[str | None, str | None]] = field(default_factory=dict)
    by_internal: dict[str, tuple[str | None, str | None]] = field(default_factory=dict)

    def lookup(self, marketplace_sku: str | None,
               internal_sku: str) -> tuple[str | None, str | None]:
        """``(barcode, description)``; either may be None.

        Marketplace SKU first (that's what the label is printed for), internal
        SKU as fallback for packs that have no marketplace listing yet. The
        description is returned even without a barcode: the two are
        independent, and the product name is almost always known.
        """
        for key, index in ((marketplace_sku, self.by_marketplace),
                           (internal_sku, self.by_internal)):
            if key and key in index:
                return index[key]
        return (None, None)


def describe(label: str, sub: str) -> str | None:
    """``"Label - Sub"``, without a dangling dash when Sub is empty.

    The old spreadsheet always appended " - ", which is why the pre-send
    checklist had an item to remove trailing dashes by hand.
    """
    if label and sub:
        return f"{label} - {sub}"
    return label or sub or None


def load_barcodes(path: Path) -> BarcodeCatalog:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        catalog = BarcodeCatalog()
        for row in wb[BARCODE_SHEET].iter_rows(min_row=2, max_col=8, values_only=True):
            barcode = (_clean(_at(row, BARCODE["barcode"]))
                       or _clean(_at(row, BARCODE["barcode2"])) or None)
            entry = (barcode, describe(_clean(_at(row, BARCODE["label"])),
                                       _clean(_at(row, BARCODE["sub"]))))
            for col, index in ((BARCODE["marketplace_sku"], catalog.by_marketplace),
                               (BARCODE["internal_sku"], catalog.by_internal)):
                key = _clean(_at(row, col))
                if key:
                    index.setdefault(key, entry)
        return catalog
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Case dimensions (ERP catalog export)
# ---------------------------------------------------------------------------

_DIMS = re.compile(r"(\d+(?:[.,]\d+)?)\s*[xX*]\s*(\d+(?:[.,]\d+)?)\s*[xX*]\s*(\d+(?:[.,]\d+)?)")
CONFIRMED = "confirmed"
_EMPTY_NOTES = {"", "-", "n/a", "n/a (not applicable)", "na", "none"}

Dims = tuple[float, float, float]


def parse_case_dimensions(text: object) -> tuple[str, Dims | None]:
    """Classify a free-text case-dimension cell.

    Returns ``("confirmed", (L, W, H))``, ``("unconfirmed", None)`` or
    ``("missing", None)``.

    Why only "Confirmed" values are used: the field was typed by hand for
    years and mixes two generations of data::

        "470 x 270 x 315 Confirmed 2041"   measured and confirmed (millimeters)
        "45 x 32 x 34"                     legacy, never verified (centimeters)

    In the production data the two groups lined up exactly with the two units.
    The CBM formula assumes millimeters, so letting a centimeter value through
    would make the CBM 1,000x too small and the vendor would quote freight on
    an absurd number. Leaving the cell blank (and asking) is the safer failure.
    """
    raw = str(text or "").strip()
    match = _DIMS.search(raw) if raw else None
    if not match:
        return "missing", None
    if CONFIRMED not in raw.lower():
        return "unconfirmed", None
    dims = tuple(round(float(x.replace(",", ".")), 1) for x in match.groups())
    return "confirmed", dims  # type: ignore[return-value]


@dataclass
class CaseDimensions:
    measures: dict[str, Dims] = field(default_factory=dict)
    notes: dict[str, str] = field(default_factory=dict)
    unconfirmed: list[tuple[str, str]] = field(default_factory=list)  # (sku, raw text)
    product_ids: set[str] = field(default_factory=set)


def _note(value: object) -> str:
    text = str(value or "").strip()
    return "" if text.lower() in _EMPTY_NOTES else text


def load_case_dimensions(path: Path) -> CaseDimensions:
    """Read the ERP catalog export (one sheet, columns located by header name)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        header = {str(v or "").strip(): i for i, v in enumerate(next(rows, ()))}
        i_sku = header.get("ProductID")
        if i_sku is None:
            raise ValueError(f"{path.name}: dimensions export has no ProductID column")
        i_dims = header.get("SUPPLIERCASEDIMENSIONS")
        i_note1 = header.get("PACKINGINSTRUCTIONS")
        i_note2 = header.get("PACKINGINSTRUCTIONS2")

        result = CaseDimensions()
        for row in rows:
            sku = _clean(_at(row, i_sku))
            if not sku:
                continue
            result.product_ids.add(sku)
            raw = _at(row, i_dims)
            status, dims = parse_case_dimensions(raw)
            if status == "confirmed" and dims:
                result.measures[sku] = dims
            elif status == "unconfirmed":
                result.unconfirmed.append((sku, str(raw).strip()))
            note = " / ".join(t for t in (_note(_at(row, i_note1)), _note(_at(row, i_note2))) if t)
            if note:
                result.notes[sku] = note
        log.info("Dimensions: %d confirmed, %d unconfirmed (discarded)",
                 len(result.measures), len(result.unconfirmed))
        return result
    finally:
        wb.close()


# ---------------------------------------------------------------------------
# Sales report priority
# ---------------------------------------------------------------------------

def load_priorities(path: Path) -> dict[str, int | str]:
    """``{SKU: priority}`` from the sales report, columns found by header.

    Written into the vendor file as a *value*, so the file sent by email does
    not depend on the sales report sitting next to it.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[SALES_SHEET].iter_rows(values_only=True)
        header = [normalize(v) for v in next(rows, ())]
        try:
            i_sku = header.index(SALES_SKU_HEADER)
            i_pri = next(i for i, h in enumerate(header) if h.startswith(SALES_PRIORITY_HEADER))
        except (ValueError, StopIteration) as exc:
            raise ValueError(f"{path.name}: sales report lacks SKU/priority columns") from exc
        priorities: dict[str, int | str] = {}
        for row in rows:
            sku = _clean(_at(row, i_sku))
            value = _at(row, i_pri)
            if not sku or sku in priorities or value in (None, ""):
                continue
            try:
                priorities[sku] = int(round(float(value)))
            except (TypeError, ValueError):
                priorities[sku] = str(value).strip()
        return priorities
    finally:
        wb.close()
