"""Where the case-dimensions export comes from: a folder, or the ERP via a browser.

``ExportSource`` is the seam. The build step only needs a path to a catalog
export; whether it was dropped in ``Inbox/`` or downloaded a minute ago by a
browser session is not its concern.

- ``LocalFolderSource`` picks the newest dimensions export in a folder. Used
  by the demo and the tests, and as the manual fallback.
- ``PlaywrightExportSource`` shows the production pattern for ERPs with no
  API for this export (e.g. SellerCloud-style catalog screens): a persistent
  browser profile so 2FA is needed once every few weeks, a pause for the
  operator to log in, a check that no stray filter is applied, a "select all
  N items" step whose count is *verified*, a check that the export template
  carries the right columns, and a diagnostic screenshot + HTML dump on any
  failure. Every selector is configuration; nothing is tied to a real URL.

Why the select-all count is verified twice
------------------------------------------
Grid UIs select only the visible page when you tick the header checkbox; a
separate "Select all N items" link selects the rest. If that click is missed,
the export downloads perfectly well-formed with 50 of 273 SKUs, and the vendor
file comes out with most boxes missing dimensions and nothing complaining. So
the selected counter is checked on the page, and the downloaded file is
checked again against the SKUs that were requested.
"""

from __future__ import annotations

import logging
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from .columns import DIMENSIONS_COLUMNS
from .inputs import classify

log = logging.getLogger(__name__)


class ExportError(Exception):
    """The export could not be produced or is incomplete."""


class ExportSource(ABC):
    @abstractmethod
    def fetch_case_dimensions(self, skus: list[str]) -> Path:
        """Return the path of a catalog export covering ``skus``."""


class LocalFolderSource(ExportSource):
    """Newest dimensions export found (by content) in a folder."""

    def __init__(self, folder: Path) -> None:
        self.folder = folder

    def fetch_case_dimensions(self, skus: list[str]) -> Path:
        path = classify(self.folder).get("dimensions")
        if path is None:
            raise ExportError(f"no case-dimensions export in {self.folder}")
        missing = coverage_gaps(path, skus)
        if missing:
            log.warning("%d requested SKUs are not in %s (e.g. %s)",
                        len(missing), path.name, ", ".join(missing[:5]))
        return path


def coverage_gaps(path: Path, skus: list[str]) -> list[str]:
    """SKUs requested but absent from the downloaded file."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb[wb.sheetnames[0]].iter_rows(values_only=True)
        header = [str(v or "").strip() for v in next(rows, ())]
        i = header.index("ProductID")
        present = {str(r[i]).strip() for r in rows if i < len(r) and r[i]}
    finally:
        wb.close()
    return [s for s in skus if s not in present]


@dataclass
class BrowserConfig:
    """Selectors and URLs for a catalog-export screen. All values are examples."""

    search_url: str                                   # e.g. "https://erp.example.com/catalog?skus={skus}"
    login_check_selector: str = "nav .user-menu"
    status_filter_chip: str = ".filter-chip:has-text('Active')"
    header_checkbox: str = "thead input[type=checkbox]"
    select_all_link: str = "a:has-text('Select all')"
    selected_counter: str = ".selection-count"        # text like "273 of 273 selected"
    actions_button: str = "button.bulk-actions"
    export_menu_item: str = "text=Export Products"
    custom_radio: str = "label:has-text('Custom')"
    next_button: str = "button:has-text('NEXT')"
    template_select: str = "select#Templates"
    template_name: str = "Packing Instructions"
    export_button: str = "button:has-text('EXPORT')"
    job_link: str = "a.job-link"
    download_button: str = "a:has-text('Download File')"
    profile_dir: Path = Path(".browser-profile")
    diagnostics_dir: Path = Path("diagnostics")
    download_dir: Path = Path("Inbox")
    required_columns: tuple[str, ...] = field(default=DIMENSIONS_COLUMNS)
    job_timeout_s: int = 600


class PlaywrightExportSource(ExportSource):
    """Download the catalog export through the ERP's web UI (optional extra)."""

    def __init__(self, config: BrowserConfig, headless: bool = False) -> None:
        self.cfg = config
        self.headless = headless

    def fetch_case_dimensions(self, skus: list[str]) -> Path:
        try:
            from playwright.sync_api import sync_playwright
        except ImportError as exc:  # pragma: no cover - optional dependency
            raise ExportError("install the browser extra: pip install -e .[browser]") from exc

        cfg = self.cfg
        with sync_playwright() as pw:
            context = pw.chromium.launch_persistent_context(
                str(cfg.profile_dir), headless=self.headless, accept_downloads=True)
            page = context.pages[0] if context.pages else context.new_page()
            try:
                page.goto(cfg.search_url.format(skus=",".join(skus)))
                self._ensure_logged_in(page)
                if page.locator(cfg.status_filter_chip).count():
                    raise ExportError("a status filter is applied; the export would skip SKUs")
                self._select_all(page, expected=len(skus))
                page.click(cfg.actions_button)
                page.click(cfg.export_menu_item)
                page.click(cfg.custom_radio)
                page.click(cfg.next_button)
                self._choose_template(page)
                path = self._export_and_download(page)
            except Exception as exc:
                self._dump_diagnostics(page, "export-failed")
                raise ExportError(f"browser export failed: {exc}") from exc
            finally:
                context.close()

        missing = coverage_gaps(path, skus)
        if missing:
            raise ExportError(f"export is missing {len(missing)} of {len(skus)} SKUs "
                              f"(e.g. {', '.join(missing[:5])}); 'select all' likely missed")
        return path

    def _ensure_logged_in(self, page) -> None:
        """Session lives in the persistent profile; pause for 2FA only when expired."""
        if page.locator(self.cfg.login_check_selector).count():
            return
        input("  Log in (including 2FA) in the browser window, then press Enter... ")
        page.goto(page.url)
        if not page.locator(self.cfg.login_check_selector).count():
            raise ExportError("still not logged in")

    def _select_all(self, page, expected: int) -> None:
        cfg = self.cfg
        page.click(cfg.header_checkbox)            # selects the visible page only
        if page.locator(cfg.select_all_link).count():
            page.click(cfg.select_all_link)        # selects every result
        text = page.locator(cfg.selected_counter).inner_text()
        numbers = [int(n.replace(",", "")) for n in re.findall(r"\d[\d,]*", text)]
        if not numbers or numbers[0] < expected:
            raise ExportError(f"selected {numbers[0] if numbers else 0} of {expected} items")
        log.info("Selected %d of %d items", numbers[0], expected)

    def _choose_template(self, page) -> None:
        """Load the export template, then verify its columns *before* exporting.

        A wrong template still exports a well-formed file -- with the wrong
        columns, and a vendor file with no dimensions.
        """
        page.select_option(self.cfg.template_select, label=self.cfg.template_name)
        page.wait_for_timeout(2_000)
        text = page.inner_text("body")
        missing = [c for c in self.cfg.required_columns if c not in text]
        if missing:
            raise ExportError(f"template {self.cfg.template_name!r} lacks {', '.join(missing)}")

    def _export_and_download(self, page) -> Path:
        cfg = self.cfg
        page.click(cfg.export_button)               # queues a background job
        page.click(cfg.job_link)
        page.wait_for_selector(cfg.download_button, timeout=cfg.job_timeout_s * 1000)
        with page.expect_download() as info:
            page.click(cfg.download_button)
        download = info.value
        cfg.download_dir.mkdir(parents=True, exist_ok=True)
        target = cfg.download_dir / f"Catalog Dimensions {download.suggested_filename}"
        download.save_as(target)
        return target

    def _dump_diagnostics(self, page, label: str) -> None:
        """Screenshot + HTML of the failing page: selectors break silently otherwise."""
        try:
            self.cfg.diagnostics_dir.mkdir(parents=True, exist_ok=True)
            page.screenshot(path=str(self.cfg.diagnostics_dir / f"{label}.png"), full_page=True)
            (self.cfg.diagnostics_dir / f"{label}.html").write_text(page.content(), encoding="utf8")
        except Exception as exc:  # diagnostics must never mask the original error
            log.warning("could not write diagnostics: %s", exc)
