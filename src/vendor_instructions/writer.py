"""Write the vendor workbook on top of the template, without breaking it.

The three tabs that change are rewritten; everything else (instructions,
label samples, the pre-send checklist) is kept as-is.

Decisions made here
-------------------
1. **Cross-file lookups become values; in-sheet math stays as formulas.**
   The template's lookups pointed at *other workbooks*. Emailed to a vendor,
   those turn into ``#REF!`` or, worse, stale cached values nobody can audit.
   So anything that came from another file (marketplace SKU, description,
   barcode, priority, dimensions) is written as a value. Anything computed
   inside the sheet stays a live formula::

       E  invoice total per item  SUMIF over column D
       J  PCS/CASE                =I*H
       K  # OF BAGS               =D/H
       L  # OF CASES              =K/I
       U  CBM                     =R*S*T/1e9  (dimensions in mm)
       W  G.W(KGS)                =V*L        (V is typed by the vendor)
       X  MEAS(M3)                =U*L

   K and L duplicate numbers the plan already has, deliberately: if anyone
   corrects a total in D, everything downstream recalculates instead of
   silently lying. The vendor fills in V, so W and X must recalc on *their*
   machine -- values would leave the file dead.
2. **Fragile formulas are replaced with direct references.** Labels D/E used
   to look up Box Information *by marketplace SKU*, which breaks exactly on
   the rows that have no marketplace SKU (the ones already in trouble). They
   now point straight at the row (``='Box Information'!K42``).
3. **Formatting is cloned from a template row**, cell by cell, so borders,
   fonts and number formats match the template exactly.
"""

from __future__ import annotations

import logging
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import media, template
from .assemble import BuildResult
from .columns import (
    BOX,
    BOX_LAST_COL,
    BOX_SHEET,
    INVOICE_HEADER,
    INVOICE_SHEET,
    LABELS,
    LABELS_LAST_COL,
    LABELS_SHEET,
)

log = logging.getLogger(__name__)

NO_LABEL = "DO NOT PRINT LABEL"
RED = PatternFill("solid", fgColor="FFFF7C80")      # missing and blocking
AMBER = PatternFill("solid", fgColor="FFFFC000")    # missing data, not wrong
GREEN = PatternFill("solid", fgColor="FFC6EFCE")    # checked and OK
_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def col(key: str) -> str:
    """Box Information column letter: ``col("total_pcs") -> "D"``.

    Box Information only: it is the one sheet formulas point at, and BOX and
    LABELS share key names that mean different columns.
    """
    return get_column_letter(BOX[key])


def write_workbook(template_path: Path, output: Path, result: BuildResult,
                   po: int | None, vendor: str, have_dims: bool) -> list[str]:
    """Write the vendor file. Returns warnings (empty list = clean)."""
    wb = openpyxl.load_workbook(template_path)
    _write_box_information(wb[BOX_SHEET], result, po, have_dims)
    _write_product_labels(wb[LABELS_SHEET], result)
    _write_invoice_list(wb[INVOICE_SHEET], result, vendor, po)

    # The cross-file formulas are gone, so the external links must go too.
    wb._external_links = []
    # Images are restored byte-for-byte by media.restore(); letting openpyxl
    # also write (or re-encode) them would produce duplicates or altered bytes.
    for ws in wb.worksheets:
        ws._images = []
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)

    warnings = []
    try:
        if not media.restore(template_path, output):
            warnings.append("template sample images could not be restored")
    except (OSError, KeyError, ValueError) as exc:
        warnings.append(f"template sample images could not be restored ({exc})")
    log.info("Wrote %s", output)
    return warnings


# ---------------------------------------------------------------------------
# Box Information
# ---------------------------------------------------------------------------

def _invoice_total_formula(row: int) -> str:
    """Product total shown only on the first row of each item (the parent row).

    This is the template's formula with an off-by-one fixed: the original's
    first row compared against the row *below* instead of the row above.
    """
    a, d = col("item_sku"), col("total_pcs")
    return (f'=IF(${a}{row}="","",IF(${a}{row}=${a}{row - 1},"",'
            f"SUMIF(${a}:${a},${a}{row},${d}:${d})))")


def _write_box_information(ws: Worksheet, result: BuildResult, po: int | None,
                           have_dims: bool) -> None:
    table = template.box_table(ws)
    style, height = _row_style(ws, table.style_row, BOX_LAST_COL)
    _clear(ws, table.first_data, BOX_LAST_COL)

    row = table.first_data
    for group in result.groups:
        _apply(ws, row, style, height)
        ws.cell(row, BOX["item_sku"], group.parent)
        ws.cell(row, BOX["invoice_total"], _invoice_total_formula(row))
        group.box_row = row
        row += 1

        for pack in group.packs:
            _apply(ws, row, style, height)
            pack.box_row = row
            ws.cell(row, BOX["item_sku"], group.parent)
            ws.cell(row, BOX["invoice_total"], _invoice_total_formula(row))
            _value_or_fill(ws.cell(row, BOX["fba_sku"]), pack.fba_sku, RED)
            _value_or_fill(ws.cell(row, BOX["description"]), pack.description, RED)
            ws.cell(row, BOX["total_pcs"], _int(pack.total))
            ws.cell(row, BOX["pcs_bag"], _int(pack.pcs_per_bag))
            ws.cell(row, BOX["bags_case"], _int(pack.bags_per_case))
            ws.cell(row, BOX["pcs_case"], f"={col('bags_case')}{row}*{col('pcs_bag')}{row}")
            ws.cell(row, BOX["n_bags"], f"={col('total_pcs')}{row}/{col('pcs_bag')}{row}")
            ws.cell(row, BOX["n_cases"], f"={col('n_bags')}{row}/{col('bags_case')}{row}")
            if po is not None:
                ws.cell(row, BOX["po"], po)

            if pack.dims:
                keys = ("length_mm", "width_mm", "height_mm")
                for key, value in zip(keys, pack.dims, strict=True):
                    ws.cell(row, BOX[key], _int(value))
                ws.cell(row, BOX["cbm"], f"={col('length_mm')}{row}*{col('width_mm')}{row}"
                                         f"*{col('height_mm')}{row}/1000000000")
                ws.cell(row, BOX["gross_weight"],
                        f"={col('net_weight')}{row}*{col('n_cases')}{row}")
                ws.cell(row, BOX["measurement"], f"={col('cbm')}{row}*{col('n_cases')}{row}")
            elif have_dims:
                # The export exists but this SKU has no confirmed size: amber
                # means "data missing", not "data wrong".
                for key in ("length_mm", "width_mm", "height_mm"):
                    ws.cell(row, BOX[key]).fill = AMBER
            if pack.note:
                ws.cell(row, BOX["notes"], pack.note)
            row += 1

    _fit_widths(ws, table.first_data, row - 1,
                (BOX["item_sku"], BOX["fba_sku"], BOX["description"], BOX["notes"]))
    ws.freeze_panes = ws.cell(table.first_data, BOX["fba_sku"])


# ---------------------------------------------------------------------------
# Product Labels
# ---------------------------------------------------------------------------

def _write_product_labels(ws: Worksheet, result: BuildResult) -> None:
    table = template.labels_table(ws)
    style, height = _row_style(ws, table.style_row, LABELS_LAST_COL)
    _clear(ws, table.first_data, LABELS_LAST_COL)

    row = table.first_data
    for group in result.groups:
        _apply(ws, row, style, height)
        ws.cell(row, LABELS["item_sku"], group.parent)
        _value_or_fill(ws.cell(row, LABELS["priority"]), group.priority, AMBER)
        ws.cell(row, LABELS["total"], f"='{BOX_SHEET}'!{col('invoice_total')}{group.box_row}")
        row += 1

        for pack in group.packs:
            _apply(ws, row, style, height)
            ws.cell(row, LABELS["pack_size"], _int(pack.pcs_per_bag))
            ws.cell(row, LABELS["packs_req"], f"='{BOX_SHEET}'!{col('n_bags')}{pack.box_row}")
            _value_or_fill(ws.cell(row, LABELS["sku"]), pack.fba_sku, RED)
            barcode = ws.cell(row, LABELS["barcode"])
            if pack.barcode:
                barcode.value = pack.barcode
            else:
                # No barcode, no label. Keep the template's font, just bold it.
                barcode.value = NO_LABEL
                barcode.fill = RED
                font = copy(barcode.font)
                font.b = True
                barcode.font = font
            if pack.description:
                ws.cell(row, LABELS["description"], pack.description)
            row += 1

    _fit_widths(ws, table.first_data, row - 1,
                (LABELS["sku"], LABELS["barcode"], LABELS["description"]))


# ---------------------------------------------------------------------------
# Packing List & Invoice Info
# ---------------------------------------------------------------------------

def _write_invoice_list(ws: Worksheet, result: BuildResult, vendor: str, po: int | None) -> None:
    """One line per product (not per pack) for the commercial invoice.

    Totals are SUMIFs against Box Information, so a correction there flows
    through. The template's example list is wiped first (see template.py).
    """
    start = template.invoice_example_row(ws)
    _clear(ws, start, max(3, ws.max_column))

    title = f"TOTAL ORDERED PER PRODUCT - {vendor}" + (f"  |  PO #{po}" if po is not None else "")
    ws.cell(start + 1, 1, title).font = Font(b=True, sz=14)
    header = start + 3
    for c, text in enumerate(INVOICE_HEADER, 1):
        cell = ws.cell(header, c, text)
        cell.font = Font(b=True, sz=12)
        cell.fill = PatternFill("solid", fgColor="FFFFFF00")
        cell.border = _BORDER

    a, d = col("item_sku"), col("total_pcs")
    row = first = header + 1
    for group in result.groups:
        ws.cell(row, 1, group.parent)
        ws.cell(row, 2, group.description)
        total = ws.cell(row, 3, f"=SUMIF('{BOX_SHEET}'!${a}:${a},$A{row},"
                                f"'{BOX_SHEET}'!${d}:${d})")
        total.number_format = "#,##0"
        for c in (1, 2, 3):
            ws.cell(row, c).border = _BORDER
        row += 1

    ws.cell(row, 1, "TOTAL").font = Font(b=True)
    grand = ws.cell(row, 3, f"=SUM(C{first}:C{row - 1})")
    grand.font = Font(b=True)
    grand.number_format = "#,##0"
    for c in (1, 2, 3):
        ws.cell(row, c).border = _BORDER
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 16
    ws.freeze_panes = f"A{header + 1}"


# ---------------------------------------------------------------------------
# Formatting helpers
# ---------------------------------------------------------------------------

def _row_style(ws: Worksheet, row: int, last_col: int) -> tuple[list, float | None]:
    return ([copy(ws.cell(row, c)._style) for c in range(1, last_col + 1)],
            ws.row_dimensions[row].height)


def _apply(ws: Worksheet, row: int, style: list, height: float | None) -> None:
    for c, s in enumerate(style, 1):
        ws.cell(row, c)._style = copy(s)
    if height:
        ws.row_dimensions[row].height = height


def _clear(ws: Worksheet, start: int, last_col: int) -> None:
    """Wipe old template content from ``start`` down (values, fills, comments).

    Comments go too: they are anchored to a cell, and after rewriting, that
    cell is a different product. They are also internal notes, and this file
    is sent to an outside vendor.
    """
    for row in range(start, ws.max_row + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row, c)
            cell.value = None
            if cell.fill.patternType:
                cell.fill = PatternFill()
            cell.comment = None


def _value_or_fill(cell, value, fill: PatternFill) -> None:
    if value not in (None, ""):
        cell.value = value
    else:
        cell.fill = fill


def _fit_widths(ws: Worksheet, first: int, last: int, columns: tuple[int, ...]) -> None:
    """Widen text columns so nothing is cut off -- only widen, and cap at 55."""
    for c in columns:
        longest = max((len(v) for r in range(first, last + 1)
                       if isinstance(v := ws.cell(r, c).value, str) and not v.startswith("=")),
                      default=0)
        if longest:
            letter = get_column_letter(c)
            current = ws.column_dimensions[letter].width or 0
            ws.column_dimensions[letter].width = max(current, min(longest + 3, 55))


def _int(value: float | None) -> float | int | None:
    """Whole numbers as int, so Excel shows 450 instead of 450.0."""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value
