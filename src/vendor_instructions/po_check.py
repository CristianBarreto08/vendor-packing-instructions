"""`verify`: double-check a finished instructions file against the ERP PO export.

Between generating the instructions and sending them, the file gets touched:
quantities are corrected by hand, a product is added after a call with the
vendor, a value gets pasted over a formula. The purchase order in the ERP is
what will be invoiced and received against, so before the file goes out it is
reconciled against the PO export line by line:

- **Quantities**: per item SKU, sum of TOTAL PIECES == PO quantity ordered.
- **SKU coverage**: every PO line is packed; nothing is packed that is not on
  the PO.
- **PO number**: every pack row carries the PO number of the export.
- **Pack math**: totals are whole bags; bags fill whole cases (a short last
  carton is a warning, not an error).
- **Formula integrity**: PCS/CASE, # OF BAGS and # OF CASES are still the live
  formulas -- or, if someone pasted values, the values still agree.
- **Labels**: packs with no FBA SKU, or flagged DO NOT PRINT LABEL.

It writes a ``CHECKED`` copy (cells colored green/amber/red plus a ``PO Check``
tab) and a notes file, and never modifies the original.
"""

from __future__ import annotations

import logging
import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

from . import media, template
from .columns import BOX, BOX_SHEET, LABELS, LABELS_SHEET, PO_EXPORT_ALIASES, normalize
from .inputs import po_from_filename
from .writer import AMBER, GREEN, NO_LABEL, RED, col

log = logging.getLogger(__name__)

ERROR, WARNING = "error", "warning"


@dataclass
class Finding:
    severity: str
    sku: str
    check: str
    message: str
    cell: str | None = None


@dataclass
class PackRow:
    row: int
    item_sku: str
    fba_sku: str | None
    total: float | None
    pcs_per_bag: float | None
    bags_per_case: float | None
    po: object


@dataclass
class PurchaseOrder:
    number: int | None
    quantities: dict[str, float]
    names: dict[str, str] = field(default_factory=dict)


@dataclass
class CheckReport:
    po_number: int | None
    lines: list[tuple[str, float | None, float | None]]   # sku, PO qty, packed qty
    findings: list[Finding]

    @property
    def errors(self) -> list[Finding]:
        return [f for f in self.findings if f.severity == ERROR]

    @property
    def warnings(self) -> list[Finding]:
        return [f for f in self.findings if f.severity == WARNING]


# ---------------------------------------------------------------------------
# Reading
# ---------------------------------------------------------------------------

def read_po_export(path: Path) -> PurchaseOrder:
    """Read a PO export; the header row and columns are found by alias."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    finally:
        wb.close()
    located = _locate_header(rows)
    if located is None:
        raise ValueError(f"{path.name}: no header row with a SKU and a quantity column")
    header_row, idx = located
    width = max(i for i in idx.values() if i is not None) + 1

    quantities: dict[str, float] = defaultdict(float)
    names: dict[str, str] = {}
    numbers: set[int] = set()
    for raw in rows[header_row + 1:]:
        raw = tuple(raw) + (None,) * (width - len(raw))
        sku = str(raw[idx["sku"]] or "").strip()
        if not sku:
            continue
        quantities[sku] += _num(raw[idx["qty"]]) or 0.0   # split lines are summed
        if idx["name"] is not None and raw[idx["name"]]:
            names.setdefault(sku, str(raw[idx["name"]]).strip())
        if idx["po"] is not None and (n := _num(raw[idx["po"]])) is not None:
            numbers.add(int(n))
    if len(numbers) > 1:
        raise ValueError(f"{path.name}: export mixes several POs ({sorted(numbers)})")
    number = numbers.pop() if numbers else po_from_filename(path)
    return PurchaseOrder(number, dict(quantities), names)


def _locate_header(rows: list[tuple]) -> tuple[int, dict[str, int | None]] | None:
    """Header row (exports often carry a title line first) and column indexes."""
    for n, raw in enumerate(rows[:15]):
        header = [normalize(v) for v in raw]
        idx = {key: next((header.index(a) for a in aliases if a in header), None)
               for key, aliases in PO_EXPORT_ALIASES.items()}
        if idx["sku"] is not None and idx["qty"] is not None:
            return n, idx
    return None


def read_pack_rows(ws: Worksheet) -> list[PackRow]:
    """Pack rows of Box Information (rows with a TOTAL PIECES value)."""
    table = template.box_table(ws)
    rows = []
    for r in range(table.first_data, ws.max_row + 1):
        item = ws.cell(r, BOX["item_sku"]).value
        total = ws.cell(r, BOX["total_pcs"]).value
        if not item or total in (None, ""):
            continue  # blank or parent row
        rows.append(PackRow(
            row=r, item_sku=str(item).strip(),
            fba_sku=(str(v).strip() if (v := ws.cell(r, BOX["fba_sku"]).value) else None),
            total=_num(total), pcs_per_bag=_num(ws.cell(r, BOX["pcs_bag"]).value),
            bags_per_case=_num(ws.cell(r, BOX["bags_case"]).value),
            po=ws.cell(r, BOX["po"]).value))
    return rows


def _num(value: object) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return None


# ---------------------------------------------------------------------------
# Checks
# ---------------------------------------------------------------------------

def check(ws_box: Worksheet, ws_labels: Worksheet | None, po: PurchaseOrder) -> CheckReport:
    packs = read_pack_rows(ws_box)
    findings: list[Finding] = []

    packed: dict[str, float] = defaultdict(float)
    for p in packs:
        packed[p.item_sku] += p.total or 0.0
        findings += _pack_math(p)
        findings += _formulas_intact(ws_box, p)
        findings += _po_number(p, po.number)
        if not p.fba_sku:
            findings.append(Finding(WARNING, p.item_sku, "label",
                                    "pack has no FBA SKU: the vendor cannot label it",
                                    f"{col('fba_sku')}{p.row}"))

    lines = []
    for sku in list(dict.fromkeys([*po.quantities, *packed])):
        ordered, in_file = po.quantities.get(sku), packed.get(sku)
        lines.append((sku, ordered, in_file))
        if in_file is None:
            findings.append(Finding(ERROR, sku, "coverage",
                                    f"on the PO ({ordered:,.0f} pcs) but not in the instructions"))
        elif ordered is None:
            findings.append(Finding(ERROR, sku, "coverage",
                                    f"packed ({in_file:,.0f} pcs) but not on the PO"))
        elif abs(ordered - in_file) > 0.5:
            findings.append(Finding(ERROR, sku, "quantity",
                                    f"PO says {ordered:,.0f} pcs, instructions pack "
                                    f"{in_file:,.0f} ({in_file - ordered:+,.0f})"))

    if ws_labels is not None:
        findings += _no_label_rows(ws_labels)
    findings.sort(key=lambda f: f.severity != ERROR)  # errors first, order otherwise kept
    return CheckReport(po.number, lines, findings)


def _pack_math(p: PackRow) -> list[Finding]:
    where = f"{col('total_pcs')}{p.row}"
    if not p.total or not p.pcs_per_bag or not p.bags_per_case:
        return [Finding(ERROR, p.item_sku, "pack math",
                        "total, PCS/BAG or BAGS/CASE is blank or zero", where)]
    bags = p.total / p.pcs_per_bag
    if abs(bags - round(bags)) > 1e-9:
        return [Finding(ERROR, p.item_sku, "pack math",
                        f"{p.total:,.0f} pcs is not a whole number of {p.pcs_per_bag:g}-pc bags",
                        where)]
    full, rest = divmod(round(bags), int(p.bags_per_case))
    if rest:
        return [Finding(WARNING, p.item_sku, "pack math",
                        f"{round(bags):,} bags = {full} full cases + a short case of {rest} bags",
                        f"{col('n_cases')}{p.row}")]
    return []


def _formulas_intact(ws: Worksheet, p: PackRow) -> list[Finding]:
    """Live formulas expected; pasted values are tolerated only if still correct."""
    r = p.row
    expected = {
        "pcs_case": (f"={col('bags_case')}{r}*{col('pcs_bag')}{r}",
                     (p.bags_per_case or 0) * (p.pcs_per_bag or 0)),
        "n_bags": (f"={col('total_pcs')}{r}/{col('pcs_bag')}{r}",
                   (p.total or 0) / p.pcs_per_bag if p.pcs_per_bag else None),
    }
    if expected["n_bags"][1] is not None and p.bags_per_case:
        expected["n_cases"] = (f"={col('n_bags')}{r}/{col('bags_case')}{r}",
                               expected["n_bags"][1] / p.bags_per_case)
    findings = []
    for key, (formula, value) in expected.items():
        cell = ws.cell(r, BOX[key])
        content = cell.value
        if isinstance(content, str) and content.startswith("="):
            if re.sub(r"[\s$]", "", content).upper() != formula.upper():
                findings.append(Finding(WARNING, p.item_sku, "formula",
                                        f"unexpected formula {content}", cell.coordinate))
        elif value is not None and (_num(content) is None or abs(_num(content) - value) > 1e-6):
            findings.append(Finding(ERROR, p.item_sku, "formula",
                                    f"hard-coded value {content!r} should be {value:g}",
                                    cell.coordinate))
        elif value is not None:
            findings.append(Finding(WARNING, p.item_sku, "formula",
                                    "formula replaced by a value (still correct, but will "
                                    "not recalculate)", cell.coordinate))
    return findings


def _po_number(p: PackRow, number: int | None) -> list[Finding]:
    if number is None:
        return []
    if _num(p.po) != number:
        return [Finding(ERROR, p.item_sku, "po number",
                        f"PO # is {p.po!r}, the export is PO {number}",
                        f"{col('po')}{p.row}")]
    return []


def _no_label_rows(ws: Worksheet) -> list[Finding]:
    table = template.labels_table(ws)
    findings, parent = [], ""
    for r in range(table.first_data, ws.max_row + 1):
        parent = ws.cell(r, LABELS["item_sku"]).value or parent
        if ws.cell(r, LABELS["barcode"]).value == NO_LABEL:
            findings.append(Finding(WARNING, str(parent), "label",
                                    "flagged DO NOT PRINT LABEL: confirm it ships unlabeled",
                                    f"{LABELS_SHEET}!{ws.cell(r, LABELS['barcode']).coordinate}"))
    return findings


# ---------------------------------------------------------------------------
# Outputs
# ---------------------------------------------------------------------------

def verify(instructions: Path, po_export: Path, out_dir: Path) -> tuple[CheckReport, Path, Path]:
    """Run all checks; write the CHECKED copy and the notes file."""
    po = read_po_export(po_export)
    wb = openpyxl.load_workbook(instructions)
    labels = wb[LABELS_SHEET] if LABELS_SHEET in wb.sheetnames else None
    report = check(wb[BOX_SHEET], labels, po)

    _color_cells(wb[BOX_SHEET], report)
    _summary_sheet(wb, report)
    for ws in wb.worksheets:
        ws._images = []
    out_dir.mkdir(parents=True, exist_ok=True)
    checked = out_dir / f"{instructions.stem} - CHECKED.xlsx"
    wb.save(checked)
    media.restore(instructions, checked)

    notes = out_dir / f"{instructions.stem} - CHECK NOTES.txt"
    notes.write_text(render_notes(report, instructions, po_export), encoding="utf8")
    log.info("PO check: %d errors, %d warnings", len(report.errors), len(report.warnings))
    return report, checked, notes


def _color_cells(ws: Worksheet, report: CheckReport) -> None:
    """Green on quantities that reconcile; amber/red on each finding's cell."""
    bad_qty = {f.sku for f in report.findings if f.check in ("quantity", "coverage")}
    for p in read_pack_rows(ws):
        ws.cell(p.row, BOX["total_pcs"]).fill = RED if p.item_sku in bad_qty else GREEN
    for f in sorted(report.findings, key=lambda f: f.severity == ERROR):
        if f.cell and "!" not in f.cell:
            ws[f.cell].fill = RED if f.severity == ERROR else AMBER


def _summary_sheet(wb: openpyxl.Workbook, report: CheckReport) -> None:
    if "PO Check" in wb.sheetnames:  # re-verifying a CHECKED copy
        del wb["PO Check"]
    ws = wb.create_sheet("PO Check", 0)
    status = "PASSED" if not report.errors else f"{len(report.errors)} ERRORS"
    ws.append([f"PO {report.po_number} check: {status}, {len(report.warnings)} warnings"])
    ws["A1"].font = Font(b=True, sz=14)
    ws.append([])
    ws.append(["ITEM SKU", "PO QTY", "PACKED QTY", "DIFF", "STATUS"])
    header = ws.max_row
    for sku, ordered, packed in report.lines:
        diff = (packed or 0) - (ordered or 0)
        ok = ordered is not None and packed is not None and abs(diff) <= 0.5
        ws.append([sku, ordered, packed, diff, "OK" if ok else "MISMATCH"])
        ws.cell(ws.max_row, 5).fill = GREEN if ok else RED
    ws.append([])
    ws.append(["SEVERITY", "ITEM SKU", "CHECK", "CELL", "DETAIL"])
    findings_header = ws.max_row
    for f in report.findings:
        ws.append([f.severity.upper(), f.sku, f.check, f.cell or "", f.message])
        ws.cell(ws.max_row, 1).fill = RED if f.severity == ERROR else AMBER
    for r in (header, findings_header):
        for c in ws[r]:
            c.font = Font(b=True)
    for letter, width in zip("ABCDE", (14, 18, 14, 14, 70), strict=True):
        ws.column_dimensions[letter].width = width
    for r in range(header + 1, findings_header - 1):
        for c in (2, 3, 4):
            ws.cell(r, c).number_format = "#,##0"


def render_notes(report: CheckReport, instructions: Path, po_export: Path) -> str:
    out = ["PO CHECK NOTES", "=" * 60,
           f"instructions  {instructions.name}", f"PO export     {po_export.name}",
           f"PO number     {report.po_number}", "",
           f"Result: {'PASSED' if not report.errors else 'FAILED'} - "
           f"{len(report.errors)} errors, {len(report.warnings)} warnings", ""]
    for severity, items in (("ERRORS (fix before sending)", report.errors),
                            ("WARNINGS (review)", report.warnings)):
        out += [severity, "-" * 60]
        out += [f"  {f.sku:<14} {f.check:<11} {(f.cell or ''):<22} {f.message}"
                for f in items] or ["  none"]
        out.append("")
    return "\n".join(out)
