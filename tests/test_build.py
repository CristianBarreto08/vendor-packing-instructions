from __future__ import annotations

import openpyxl
import pytest

from vendor_instructions import media, self_check, template
from vendor_instructions.columns import BOX, BOX_SHEET, INVOICE_SHEET, LABELS, LABELS_SHEET
from vendor_instructions.pipeline import build
from vendor_instructions.writer import AMBER, NO_LABEL, RED

from .conftest import VENDOR


def _all_values(ws):
    return [c.value for row in ws.iter_rows() for c in row if c.value is not None]


def test_end_to_end_build_on_sample_data(built):
    r = built.result
    assert (len(r.groups), r.pack_count, len(r.dropped)) == (9, 18, 3)
    assert built.po == 2041
    assert built.problems == []                       # self-check passed
    assert any("ignored a dimensions export in Input/" in w for w in built.warnings)
    assert built.review.read_text().count("FBA-SC-9014-P12") == 1


def test_formulas_stay_live_and_external_links_are_gone(built):
    wb = openpyxl.load_workbook(built.output)
    box = wb[BOX_SHEET]
    pack = built.result.groups[0].packs[0]
    r = pack.box_row
    assert box.cell(r, BOX["n_bags"]).value == f"=D{r}/H{r}"
    assert box.cell(r, BOX["n_cases"]).value == f"=K{r}/I{r}"
    assert box.cell(r, BOX["cbm"]).value == f"=R{r}*S{r}*T{r}/1000000000"
    assert box.cell(r, BOX["po"]).value == 2041
    formulas = [v for ws in wb.worksheets for v in _all_values(ws)
                if isinstance(v, str) and v.startswith("=")]
    assert formulas and not any("[1]" in f or "[2]" in f for f in formulas)
    # Labels point straight at the Box Information row, not a lookup by SKU.
    labels = wb[LABELS_SHEET]
    first = template.labels_table(labels).first_data
    assert labels.cell(first + 1, LABELS["packs_req"]).value == f"='{BOX_SHEET}'!K{r}"


def test_template_leftovers_comments_and_flags(built):
    wb = openpyxl.load_workbook(built.output)
    box, labels, invoice = wb[BOX_SHEET], wb[LABELS_SHEET], wb[INVOICE_SHEET]
    # The example list (another order's SKUs) is wiped, not just overwritten.
    for ws in (box, labels, invoice):
        assert not any(isinstance(v, str) and v.startswith("PC-88") for v in _all_values(ws))
    assert not any(c.comment for row in box.iter_rows() for c in row)

    groups = {g.parent: g for g in built.result.groups}
    mf = groups["BC-5521"].packs[1]
    assert box.cell(mf.box_row, BOX["fba_sku"]).fill.fgColor.rgb == RED.fgColor.rgb
    unconfirmed = groups["SB-3305"].packs[0]
    assert box.cell(unconfirmed.box_row, BOX["length_mm"]).value is None
    assert box.cell(unconfirmed.box_row, BOX["length_mm"]).fill.fgColor.rgb == AMBER.fgColor.rgb
    no_label = [c for row in labels.iter_rows() for c in row if c.value == NO_LABEL]
    assert len(no_label) == 1 and no_label[0].font.b


def test_template_images_survive_byte_for_byte(built, workspace):
    template_path = next(workspace.reference.glob("*Template*.xlsx"))
    assert media.image_hashes(built.output) == media.image_hashes(template_path)
    assert media.anchor_count(built.output) == media.anchor_count(template_path) == 2


def test_header_is_found_after_template_rows_shift(workspace):
    path = next(workspace.reference.glob("*Template*.xlsx"))
    wb = openpyxl.load_workbook(path)
    old_header = template.box_table(wb[BOX_SHEET]).header
    wb[BOX_SHEET].insert_rows(3, amount=2)     # someone adds two instruction lines
    wb.save(path)

    outcome = build(workspace, vendor=VENDOR)
    box = openpyxl.load_workbook(outcome.output)[BOX_SHEET]
    table = template.box_table(box)
    assert table.header == old_header + 2
    assert box.cell(table.first_data, BOX["item_sku"]).value == "HX-1042"
    assert outcome.problems == []


def test_renamed_header_stops_the_run(workspace):
    path = next(workspace.reference.glob("*Template*.xlsx"))
    wb = openpyxl.load_workbook(path)
    ws = wb[BOX_SHEET]
    ws.cell(template.box_table(ws).header, BOX["fba_sku"], "MARKETPLACE SKU")
    wb.save(path)
    with pytest.raises(template.TemplateChanged):
        build(workspace, vendor=VENDOR)


def test_self_check_catches_rows_left_under_the_data(built, workspace):
    wb = openpyxl.load_workbook(built.output)
    box = wb[BOX_SHEET]
    last = template.box_table(box).first_data + built.result.row_count
    box.cell(last + 2, BOX["item_sku"], "STALE-1")
    box.cell(last + 2, BOX["fba_sku"], "=XLOOKUP(A1,[1]All!A:A,[1]All!B:B)")
    wb.save(built.output)
    template_path = next(workspace.reference.glob("*Template*.xlsx"))
    problems = self_check.check_output(built.output, built.result, 2041, template_path)
    text = " | ".join(problems)
    assert "old template rows below the data" in text
    assert "point to another workbook" in text
