from __future__ import annotations

import openpyxl

from vendor_instructions import po_check
from vendor_instructions.columns import BOX, BOX_SHEET, LABELS_SHEET


def _po_export(workspace):
    return next((workspace.root / "PO Export").glob("*.xlsx"))


def _check(path, po):
    wb = openpyxl.load_workbook(path)
    return po_check.check(wb[BOX_SHEET], wb[LABELS_SHEET], po)


def test_verify_flags_the_seeded_discrepancies(built, workspace, tmp_path):
    report, checked, notes = po_check.verify(built.output, _po_export(workspace), tmp_path)
    errors = {(f.sku, f.check) for f in report.errors}
    assert errors == {("CT-2210", "quantity"), ("KS-3307", "coverage")}
    assert any(f.check == "pack math" and "short case" in f.message for f in report.warnings)
    assert "FAILED - 2 errors" in notes.read_text()
    wb = openpyxl.load_workbook(checked)
    assert wb.sheetnames[0] == "PO Check"
    assert openpyxl.load_workbook(built.output).sheetnames[0] != "PO Check"  # original untouched


def test_po_export_lines_are_summed_and_po_taken_from_filename(tmp_path):
    path = tmp_path / "export for po 77.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Purchase order items"])            # header not on row 1
    wb.active.append(["ProductID", "ProductName", "QtyOrdered"])
    for sku, qty in (("A-1", 100), ("A-1", 50), ("B-2", 10)):
        wb.active.append([sku, "x", qty])
    wb.save(path)
    po = po_check.read_po_export(path)
    assert po.number == 77 and po.quantities == {"A-1": 150, "B-2": 10}


def test_pasted_values_partial_bags_and_wrong_po_are_caught(built, workspace):
    po = po_check.read_po_export(_po_export(workspace))
    wb = openpyxl.load_workbook(built.output)
    box = wb[BOX_SHEET]
    groups = {g.parent: g for g in built.result.groups}
    hx5, hx10 = groups["HX-1042"].packs
    box.cell(hx5.box_row, BOX["n_bags"], 999)                     # wrong pasted value
    box.cell(hx10.box_row, BOX["n_bags"], hx10.total / hx10.pcs_per_bag)  # right, but static
    box.cell(hx10.box_row, BOX["po"], 2040)                       # wrong PO on one row
    lb = groups["LB-6620"].packs[0]
    box.cell(lb.box_row, BOX["total_pcs"], lb.total + 1)          # not a whole bag
    wb.save(built.output)

    report = _check(built.output, po)
    found = {(f.check, f.cell, f.severity) for f in report.findings}
    assert ("formula", f"K{hx5.box_row}", "error") in found
    assert ("formula", f"K{hx10.box_row}", "warning") in found
    assert ("po number", f"N{hx10.box_row}", "error") in found
    assert ("pack math", f"D{lb.box_row}", "error") in found
    assert any(f.sku == "LB-6620" and f.check == "quantity" for f in report.errors)
